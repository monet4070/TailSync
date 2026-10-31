#!/usr/bin/env python3
"""Create synthetic files for the opt-in native history/preview acceptance test."""
import argparse
import hashlib
import json
import re
import shutil
import subprocess
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from reportlab.pdfgen import canvas


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    root = Path(__file__).resolve().parents[2]
    output = args.output.resolve()
    output.mkdir(parents=True, exist_ok=True)
    policy = (root / "macos/swift-ui/Sources/TailSync/Models/HistoryPreviewFileTypes.swift").read_text()
    text_extensions, image_extensions = [
        re.findall(r'"([a-z0-9]+)"', re.search(rf"{name}: Set<String> = \[(.*?)\]", policy, re.S).group(1))
        for name in ("textExtensions", "imageExtensions")
    ]
    for ext in text_extensions:
        body = f"TailSync synthetic history fixture ({ext})\n中文与 emoji 🙂\n"
        if ext in ("md", "markdown"):
            body = f"# History acceptance ({ext})\n\n**Bold** and [link](https://example.com/).\n\n- first\n- second\n"
        elif ext in ("swift", "rs", "py", "js", "ts"):
            body = f"// synthetic {ext} fixture\nfunc acceptance() {{\n    let answer = 42\n}}\n"
        (output / f"sample.{ext}").write_text(body)
    images = root / "shared/tailsync-protocol/tests/fixtures/images"
    sources = {"png": "transparent.png", "jpg": "color.jpg", "jpeg": "color.jpg",
               "gif": "animated.gif", "webp": "animated.webp", "tif": "pages.tiff", "tiff": "color.tiff"}
    for ext, source in sources.items():
        shutil.copyfile(images / source, output / f"sample.{ext}")
    for ext, image_format in (("bmp", "bmp"), ("heic", "heic"), ("heif", "heic")):
        # ImageIO is the OS format boundary; a failed encoder must fail the fixture run.
        subprocess.run(["sips", "-s", "format", image_format, str(images / "color.jpg"),
                        "--out", str(output / f"sample.{ext}")], check=True, capture_output=True, timeout=6)
    (output / "sample.svg").write_text('<svg xmlns="http://www.w3.org/2000/svg" width="100" height="80"><rect width="100" height="80" fill="#3377cc"/><text x="8" y="40">TailSync</text></svg>')
    pdf = canvas.Canvas(str(output / "sample.pdf"))
    pdf.drawString(50, 780, "TailSync PDF history acceptance")
    pdf.showPage()
    pdf.drawString(50, 780, "Page two survives storage and restore")
    pdf.save()
    doc = Document()
    doc.add_heading("TailSync DOCX acceptance", 0)
    doc.add_paragraph("Synthetic history, restore and preview fixture.")
    doc.save(output / "sample.docx")
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[0])
    slide.shapes.title.text = "TailSync PPTX acceptance"
    slides.save(output / "sample.pptx")
    workbook = Workbook()
    workbook.active["A1"] = "Metadata-only preview is expected"
    workbook.save(output / "sample.xlsx")
    native_rtf = subprocess.run(["textutil", "-convert", "rtf", "-stdin", "-stdout"],
                                input=b"History acceptance", capture_output=True, check=True, timeout=6)
    (output / "native.rtf").write_bytes(native_rtf.stdout)
    import zipfile
    with zipfile.ZipFile(output / "sample.zip", "w") as archive:
        archive.writestr("fixture.txt", "Metadata-only preview is expected")
    # The legacy PPT is an independently sourced, real document, not a signature stub.
    ppt = output / "slides.ppt"
    if not ppt.is_file():
        raise FileNotFoundError("Provide the Apache POI basic_test_ppt_file.ppt fixture as slides.ppt")
    expected = {ext: "image" for ext in image_extensions}
    expected.update({ext: "text" for ext in text_extensions})
    format_source = (root / "macos/swift-ui/Sources/TailSync/Models/HistoryPreview.swift").read_text()
    code_extensions = re.findall(r'"([a-z0-9]+)"', re.search(r"codeFileExtensions: Set<String> = \[(.*?)\]", format_source, re.S).group(1))
    expected.update({ext: "code" for ext in code_extensions if ext in text_extensions})
    expected.update({"md": "markdown", "markdown": "markdown", "svg": "svg", "pdf": "pdf",
                     "docx": "docx", "pptx": "presentation", "ppt": "presentation",
                     "xlsx": "unsupported", "zip": "unsupported"})
    files = sorted(p for p in output.iterdir() if p.name != "manifest.json" and p.suffix[1:] in expected)
    records = [{"name": p.name, "format": expected[p.suffix[1:]], "bytes": p.stat().st_size,
                "sha256": hashlib.sha256(p.read_bytes()).hexdigest()} for p in files]
    (output / "manifest.json").write_text(json.dumps(records, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({"files": len(records), "extensions": sorted(expected)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
