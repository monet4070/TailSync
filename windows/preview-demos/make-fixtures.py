#!/usr/bin/env python3
"""Generate synthetic, offline fixtures. Requires reportlab, python-docx and Pillow."""
import argparse
import hashlib
import io
import json
import math
import random
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


def illustration(seed=0, width=1200, height=600):
    image = Image.new("RGB", (width, height), "#efece4")
    draw = ImageDraw.Draw(image)
    rng = random.Random(seed)
    for i in range(80):
        x, y = rng.randrange(width), rng.randrange(height)
        size = rng.randrange(20, 180)
        draw.ellipse((x, y, x + size, y + size), fill=rng.choice(["#356554", "#8bbd9a", "#d5b779", "#b5d3d6"]))
    draw.rectangle((20, height - 65, width - 20, height - 20), fill="#faf8f2")
    draw.text((35, height - 52), "TailSync preview fixture / image " + str(seed), fill="#1c3e30")
    out = io.BytesIO(); image.save(out, format="PNG"); return out.getvalue()


def pdf(path, pages, font):
    c = canvas.Canvas(str(path), pagesize=(595, 842), invariant=1)
    c.setTitle("TailSync native preview fixture"); c.setAuthor("TailSync test suite")
    for page in range(pages):
        c.setFillColorRGB(.98, .97, .94); c.rect(0, 0, 595, 842, fill=1, stroke=0)
        c.setFillColorRGB(.12, .18, .15); c.setFont(font, 22)
        c.drawString(42, 784, "TailSync PDF preview")
        c.setFont(font, 11)
        lines = [f"Page {page + 1} of {pages} / searchable marker: ALPHA-{page + 1:03}",
                 "中文预览测试：剪贴板历史、图片、文件与表格。",
                 "Search and select this text. Numbers: 0123456789. Symbols: + - = / %."]
        for i, line in enumerate(lines): c.drawString(42, 752 - i * 22, line)
        c.drawImage(ImageReader(io.BytesIO(illustration(page % 12, 1800, 900))), 42, 420, width=510, height=255)
        for row in range(7):
            y = 385 - row * 32
            c.setFillColorRGB(.88 if row % 2 == 0 else .96, .92 if row % 2 == 0 else .96, .89 if row % 2 == 0 else .94)
            c.rect(42, y - 8, 510, 30, fill=1, stroke=0)
            c.setFillColorRGB(.12, .18, .15)
            c.drawString(52, y + 2, f"Row {row + 1}"); c.drawString(180, y + 2, "中文 English 123")
            c.drawString(420, y + 2, f"{(page + 1) * (row + 1):04}")
        c.setFont(font, 10); c.drawString(42, 91, "Expected: readable Chinese, rectangular table rows, unclipped image.")
        c.drawString(42, 60, "Generated locally. No personal documents or network resources.")
        c.showPage()
    c.save()


def normalize_zip(path):
    source = zipfile.ZipFile(path)
    entries = [(name, source.read(name)) for name in source.namelist()]; source.close()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        for name, data in entries:
            item = zipfile.ZipInfo(name, (2026, 1, 1, 0, 0, 0)); item.compress_type = zipfile.ZIP_DEFLATED
            out.writestr(item, data)


def docx(path, layout):
    d = Document(); section = d.sections[0]
    section.page_width = Inches(8.27); section.page_height = Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.75)
    d.core_properties.author = "TailSync test suite"
    d.core_properties.created = d.core_properties.modified = datetime(2026, 1, 1, tzinfo=timezone.utc)
    for name in ["Normal", "Title", "Heading 1", "Heading 2"]:
        style = d.styles[name]; style.font.name = "Arial"; style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    d.styles["Normal"].font.size = Pt(11)
    d.styles["Normal"].paragraph_format.space_after = Pt(8)
    if layout:
        section.header.paragraphs[0].text = "TailSync header / 页眉版式检查"
        footer = section.footer.paragraphs[0]; footer.text = "Page "
        field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)
    for page in range(3):
        if page: d.add_page_break()
        d.add_paragraph("TailSync document preview" if page == 0 else f"Document section {page + 1}", "Title" if page == 0 else "Heading 1")
        d.add_paragraph("This document checks native DOCX reading, text extraction, images and table layout. Compare the same file in each available engine.")
        p = d.add_paragraph("中文预览测试：剪贴板历史、图片、文件、表格与段落。")
        p.add_run(" Bold English / 粗体 ").bold = True
        p.add_run("Italic sample").italic = True
        d.add_paragraph(f"Search marker ALPHA-{page + 1:03}. Numbers 0123456789. Punctuation + - = / %.")
        table = d.add_table(rows=5, cols=3); table.style = "Table Grid"
        for row in range(5):
            for col in range(3): table.cell(row, col).text = f"{row + 1}.{col + 1} 中文 / text"
        if layout:
            table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged heading / 合并单元格"
        d.add_paragraph("Inline image / 行内图片", "Heading 2")
        d.add_picture(io.BytesIO(illustration(page)), width=Inches(5.8))
        d.add_paragraph("First list item / 第一项", "List Bullet")
        d.add_paragraph("Second list item / 第二项", "List Bullet")
    d.save(path); normalize_zip(path)


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--font", type=Path, required=True, help="TTF with Chinese glyphs; embedded as a subset in PDFs")
    p.add_argument("--output", type=Path, default=Path(__file__).parent / "fixtures")
    args = p.parse_args(); args.output.mkdir(parents=True, exist_ok=True)
    pdfmetrics.registerFont(TTFont("Fixture", str(args.font)))
    pdf(args.output / "reading-3-pages.pdf", 3, "Fixture")
    pdf(args.output / "mixed-60-pages.pdf", 60, "Fixture")
    docx(args.output / "reading.docx", False)
    docx(args.output / "layout.docx", True)
    files = []
    for file in sorted(args.output.iterdir()):
        if file.suffix in (".pdf", ".docx"):
            files.append({"name": file.name, "bytes": file.stat().st_size, "sha256": hashlib.sha256(file.read_bytes()).hexdigest()})
    (args.output / "manifest.json").write_text(json.dumps({"synthetic": True, "pdf_font": args.font.name, "files": files}, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__": main()
